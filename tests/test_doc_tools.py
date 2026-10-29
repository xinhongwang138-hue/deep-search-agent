"""多格式文档读取测试：PDF / DOCX / XLSX 与异常收敛。

不依赖真实模型与网络：样本文件在 tmp_path 现场生成，
会话工作目录用 use_workspace() 临时绑定。
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tools.file_tools import read_file_content  # noqa: E402
from tools.path_guard import use_workspace  # noqa: E402


def _make_pdf(path: Path, text: str) -> None:
    """手写最小 PDF（单页 Helvetica 文本），避免为测试引入 reportlab。"""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{idx} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1")
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode("latin-1")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode("latin-1")
    path.write_bytes(bytes(out))


def _make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("深度搜索框架测试文档")
    document.add_paragraph("第二段：用于验证段落抽取。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "药品"
    table.cell(0, 1).text = "库存"
    table.cell(1, 0).text = "复方丹参滴丸"
    table.cell(1, 1).text = "120"
    document.save(str(path))


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["region", "amount"])
    sheet.append(["华东", 1041372])
    sheet.append(["华南", 742310])
    workbook.save(str(path))


def test_pdf_extract(tmp_path: Path) -> None:
    _make_pdf(tmp_path / "sample.pdf", "Deep search agent smoke test")
    with use_workspace(tmp_path):
        result = read_file_content.invoke({"filename": "sample.pdf"})
    assert "第 1 页" in result
    assert "smoke" in result
    assert "共 1 页" in result


def test_docx_extract(tmp_path: Path) -> None:
    _make_docx(tmp_path / "sample.docx")
    with use_workspace(tmp_path):
        result = read_file_content.invoke({"filename": "sample.docx"})
    assert "深度搜索框架测试文档" in result
    assert "| 药品 | 库存 |" in result
    assert "复方丹参滴丸" in result


def test_xlsx_summary(tmp_path: Path) -> None:
    _make_xlsx(tmp_path / "sample.xlsx")
    with use_workspace(tmp_path):
        result = read_file_content.invoke({"filename": "sample.xlsx"})
    assert "工作表 销售" in result
    assert "region" in result and "amount" in result
    assert "[数值列统计]" in result
    assert "合计 1783682" in result


def test_broken_pdf_is_reported_not_raised(tmp_path: Path) -> None:
    (tmp_path / "broken.pdf").write_bytes(b"not a real pdf")
    with use_workspace(tmp_path):
        result = read_file_content.invoke({"filename": "broken.pdf"})
    assert result.startswith("PDF 解析失败")


def test_unsupported_suffix(tmp_path: Path) -> None:
    (tmp_path / "data.bin").write_bytes(b"\x00\x01")
    with use_workspace(tmp_path):
        result = read_file_content.invoke({"filename": "data.bin"})
    assert "暂不支持" in result
    assert ".pdf" in result and ".docx" in result


def test_text_file_truncated(tmp_path: Path) -> None:
    (tmp_path / "long.md").write_text("# 标题\n" + "数据" * 2000, encoding="utf-8")
    with use_workspace(tmp_path):
        result = read_file_content.invoke({"filename": "long.md", "max_chars": 1000})
    assert "内容已截断" in result
    assert len(result) < 1200


def test_outside_workspace_rejected(tmp_path: Path) -> None:
    outside = tmp_path.parent / "secret.pdf"
    outside.write_bytes(b"%PDF-1.4")
    try:
        with use_workspace(tmp_path):
            result = read_file_content.invoke({"filename": str(outside)})
        assert result.startswith("读取失败")
    finally:
        outside.unlink(missing_ok=True)
